"""
将申请书正文.md转换为Word文档
- 图片以合适尺寸插入
- 数学公式以斜体文本渲染（Word原生公式需要OLE/OMML，此处用可读文本替代）
- 表格正确渲染
- 参考文献格式化
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

PROJECT_ROOT = Path(__file__).parent.parent
MD_PATH = PROJECT_ROOT / "docs" / "申请书正文.md"
FIG_DIR = PROJECT_ROOT / "docs" / "figures"
OUT_PATH = PROJECT_ROOT / "docs" / "申请书正文.docx"


def set_cell_shading(cell, color):
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color,
    })
    shading.append(shading_elem)


def add_formatted_run(paragraph, text, bold=False, italic=False, size=None, color=None, superscript=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if superscript:
        run.font.superscript = True
    return run


def process_inline_formatting(paragraph, text, base_size=10.5):
    """处理行内格式：**粗体**、$公式$、[n]引用"""
    parts = re.split(r'(\*\*[^*]+\*\*|\$[^$]+\$|\[\d+\])', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            add_formatted_run(paragraph, part[2:-2], bold=True, size=base_size)
        elif part.startswith('$') and part.endswith('$'):
            formula = part[1:-1]
            add_formatted_run(paragraph, formula, italic=True, size=base_size, color=RGBColor(0x33, 0x33, 0x99))
        elif re.match(r'^\[\d+\]$', part):
            add_formatted_run(paragraph, part, superscript=True, size=8, color=RGBColor(0x00, 0x66, 0xCC))
        else:
            add_formatted_run(paragraph, part, size=base_size)


def add_table_from_md(doc, lines):
    """从markdown表格行创建Word表格"""
    rows_data = []
    for line in lines:
        line = line.strip()
        if re.match(r'^\|[-:\s|]+\|$', line):
            continue
        cells = [c.strip() for c in line.split('|')[1:-1]]
        if cells:
            rows_data.append(cells)

    if not rows_data:
        return

    n_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=n_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row_data in enumerate(rows_data):
        for j, cell_text in enumerate(row_data):
            if j >= n_cols:
                break
            cell = table.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            clean_text = cell_text.replace(':-:', '').replace(':-', '').replace('-:', '').strip()
            process_inline_formatting(p, clean_text, base_size=9)

            for run in p.runs:
                run.font.size = Pt(9)

            if i == 0:
                for run in p.runs:
                    run.bold = True
                set_cell_shading(cell, 'E8EAF6')

    for col_idx in range(n_cols):
        max_width = 0
        for row in table.rows:
            if col_idx < len(row.cells):
                cell_len = len(row.cells[col_idx].text)
                max_width = max(max_width, cell_len)


def add_block_formula(doc, formula_text):
    """添加块级公式（居中斜体）"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)

    formula_clean = formula_text.strip()
    formula_clean = re.sub(r'^\$\$\s*', '', formula_clean)
    formula_clean = re.sub(r'\s*\$\$$', '', formula_clean)

    run = p.add_run(formula_clean)
    run.italic = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x99)


def add_image(doc, img_path, caption=""):
    """插入图片和图注"""
    if not img_path.exists():
        p = doc.add_paragraph(f"[图片缺失: {img_path.name}]")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    run = p.add_run()
    run.add_picture(str(img_path), width=Inches(5.8))

    if caption:
        cap_p = doc.add_paragraph()
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_p.paragraph_format.space_after = Pt(12)
        clean_caption = caption.replace('**', '')
        run = cap_p.add_run(clean_caption)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def convert_md_to_docx():
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(4)

    for level in range(1, 5):
        heading_style = doc.styles[f'Heading {level}']
        heading_style.font.name = '黑体'
        heading_style.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        heading_style.font.color.rgb = RGBColor(0x1A, 0x23, 0x7E)
        if level == 1:
            heading_style.font.size = Pt(16)
        elif level == 2:
            heading_style.font.size = Pt(14)
        elif level == 3:
            heading_style.font.size = Pt(12)
        else:
            heading_style.font.size = Pt(11)

    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    i = 0
    in_block_formula = False
    block_formula_lines = []
    table_lines = []
    in_table = False
    next_caption = ""

    while i < len(lines):
        line = lines[i].rstrip('\n')
        stripped = line.strip()

        if stripped == '---':
            i += 1
            continue

        if stripped.startswith('$$') and not in_block_formula:
            content_after = stripped[2:].strip()
            if content_after.endswith('$$') and len(content_after) > 2:
                add_block_formula(doc, content_after[:-2])
            else:
                in_block_formula = True
                block_formula_lines = []
                if content_after:
                    block_formula_lines.append(content_after)
            i += 1
            continue

        if in_block_formula:
            if stripped.endswith('$$'):
                content = stripped[:-2].strip()
                if content:
                    block_formula_lines.append(content)
                full_formula = ' '.join(block_formula_lines)
                add_block_formula(doc, full_formula)
                in_block_formula = False
                block_formula_lines = []
            else:
                block_formula_lines.append(stripped)
            i += 1
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            add_table_from_md(doc, table_lines)
            in_table = False
            table_lines = []

        img_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)', stripped)
        if img_match:
            alt_text = img_match.group(1)
            img_file = img_match.group(2)
            img_path = FIG_DIR.parent / img_file
            if not img_path.exists():
                img_path = FIG_DIR / Path(img_file).name

            caption = ""
            if i + 1 < len(lines):
                next_line = lines[i + 1].strip()
                if next_line.startswith('**图') or next_line.startswith('**Fig'):
                    caption = next_line
                    i += 1

            add_image(doc, img_path, caption)
            i += 1
            continue

        if stripped.startswith('**图') and stripped.endswith('**'):
            i += 1
            continue

        if stripped.startswith('# '):
            heading = doc.add_heading(stripped[2:], level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
            i += 1
            continue

        if stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
            i += 1
            continue

        if stripped.startswith('#### '):
            doc.add_heading(stripped[5:], level=4)
            i += 1
            continue

        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            content = stripped[2:]
            process_inline_formatting(p, content, base_size=10)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        if stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            content = stripped[2:]
            process_inline_formatting(p, content, base_size=10.5)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        process_inline_formatting(p, stripped, base_size=10.5)

        i += 1

    if in_table:
        add_table_from_md(doc, table_lines)

    doc.save(str(OUT_PATH))
    print(f"[OK] Word document saved to: {OUT_PATH}")
    print(f"     File size: {OUT_PATH.stat().st_size / 1024:.0f} KB")


if __name__ == '__main__':
    convert_md_to_docx()
