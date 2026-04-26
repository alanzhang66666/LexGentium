# -*- coding: utf-8 -*-
"""将 docs/申请书正文.md 转为 Word，题目黑体四号、正文宋体五号。"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "docs" / "申请书正文.md"
OUT_PATH = ROOT / "docs" / "申请书正文.docx"

# 四号 14pt，五号 10.5pt（GB/T 9704 常用值）
PT_TITLE = Pt(14)
PT_BODY = Pt(10.5)

FONT_HEI = ("SimHei", "SimHei")
FONT_SONG = ("SimSun", "SimSun")


def _set_run_eastasia(run, font_cn: str, font_en: str, size) -> None:
    run.font.name = font_en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_cn)
    run.font.size = size


def _add_runs_with_bold(paragraph, text: str, font_cn: str, font_en: str, size) -> None:
    """将 **粗体** 拆成多个 run。"""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            inner = part[2:-2]
            run = paragraph.add_run(inner)
            _set_run_eastasia(run, font_cn, font_en, size)
            run.bold = True
        else:
            run = paragraph.add_run(part)
            _set_run_eastasia(run, font_cn, font_en, size)


def _strip_inline_math_for_word(s: str) -> str:
    """去掉 $...$ 定界符，保留可读文本（Word 中不排 LaTeX）。"""
    return re.sub(r"\$([^$]+)\$", r"\1", s)


def _prepare_paragraph_text(s: str) -> str:
    s = _strip_inline_math_for_word(s)
    return s


def _configure_normal_style(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = FONT_SONG[1]
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SONG[0])
    st.font.size = PT_BODY


def _add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_run_eastasia(run, FONT_HEI[0], FONT_HEI[1], PT_TITLE)
    p.paragraph_format.space_after = Pt(12)


def _add_section_heading(doc: Document, text: str) -> None:
    """二级标题：黑体五号加粗。"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_eastasia(run, FONT_HEI[0], FONT_HEI[1], PT_BODY)
    run.bold = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def _add_subheading(doc: Document, text: str) -> None:
    """三/四级标题：黑体五号加粗。"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_eastasia(run, FONT_HEI[0], FONT_HEI[1], PT_BODY)
    run.bold = True
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)


def _add_body_paragraph(doc: Document, text: str, italic: bool = False) -> None:
    text = _prepare_paragraph_text(text)
    p = doc.add_paragraph()
    _add_runs_with_bold(p, text, FONT_SONG[0], FONT_SONG[1], PT_BODY)
    if italic:
        for r in p.runs:
            r.italic = True
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(3)


def _add_body_no_indent(doc: Document, text: str, italic: bool = False) -> None:
    text = _prepare_paragraph_text(text)
    p = doc.add_paragraph()
    _add_runs_with_bold(p, text, FONT_SONG[0], FONT_SONG[1], PT_BODY)
    if italic:
        for r in p.runs:
            r.italic = True
    p.paragraph_format.space_after = Pt(3)


def _add_equation_block(doc: Document, lines: list[str]) -> None:
    raw = "\n".join(lines).strip()
    raw = re.sub(r"^\$\$?\s*", "", raw)
    raw = re.sub(r"\s*\$\$?$", "", raw)
    p = doc.add_paragraph()
    run = p.add_run(raw)
    _set_run_eastasia(run, FONT_SONG[0], "Times New Roman", PT_BODY)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(3)


def _add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = tbl.rows[ri].cells[ci]
            val = row[ci] if ci < len(row) else ""
            val = _prepare_paragraph_text(val)
            cell.text = val
            for p in cell.paragraphs:
                for r in p.runs:
                    _set_run_eastasia(r, FONT_SONG[0], FONT_SONG[1], PT_BODY)
    doc.add_paragraph()


def _add_image(doc: Document, md_dir: Path, alt_path: str) -> None:
    m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", alt_path.strip())
    if not m:
        return
    _, rel = m.groups()
    rel = rel.replace("/", "\\") if "\\" in str(md_dir) else rel
    img_path = (md_dir / rel).resolve()
    if not img_path.is_file():
        p = doc.add_paragraph()
        run = p.add_run(f"[图片缺失: {rel}]")
        _set_run_eastasia(run, FONT_SONG[0], FONT_SONG[1], PT_BODY)
        return
    try:
        doc.add_picture(str(img_path), width=Cm(14))
    except OSError:
        p = doc.add_paragraph()
        run = p.add_run(f"[无法插入图片: {rel}]")
        _set_run_eastasia(run, FONT_SONG[0], FONT_SONG[1], PT_BODY)
    doc.add_paragraph()


def convert(md_path: Path, out_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    md_dir = md_path.parent

    doc = Document()
    _configure_normal_style(doc)

    i = 0
    title_done = False
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if not line.strip():
            i += 1
            continue
        if re.match(r"^---+\s*$", line):
            i += 1
            continue

        if line.startswith("#### "):
            _add_subheading(doc, line[5:].strip())
            i += 1
            continue
        if line.startswith("### "):
            _add_subheading(doc, line[4:].strip())
            i += 1
            continue
        if line.startswith("## "):
            _add_section_heading(doc, line[3:].strip())
            i += 1
            continue
        if line.startswith("# ") and not line.startswith("##"):
            if not title_done:
                _add_title(doc, line[2:].strip())
                title_done = True
            else:
                _add_section_heading(doc, line[2:].strip())
            i += 1
            continue

        if line.startswith("$$"):
            if re.search(r"\$\$.+\$\$", line, re.DOTALL):
                _add_equation_block(doc, [line])
                i += 1
                continue
            block = [line]
            i += 1
            while i < len(lines):
                block.append(lines[i])
                if lines[i].rstrip().endswith("$$"):
                    i += 1
                    break
                i += 1
            _add_equation_block(doc, block)
            continue

        if line.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                l = lines[i].strip()
                if re.match(r"^\|[\s\-:|]+\|$", l.replace(" ", "")):
                    i += 1
                    continue
                cells = [c.strip() for c in l.strip("|").split("|")]
                rows.append(cells)
                i += 1
            _add_markdown_table(doc, rows)
            continue

        if line.startswith("!["):
            _add_image(doc, md_dir, line)
            i += 1
            continue

        if line.startswith("- "):
            p = doc.add_paragraph()
            _add_runs_with_bold(
                p, "• " + _prepare_paragraph_text(line[2:]), FONT_SONG[0], FONT_SONG[1], PT_BODY
            )
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        if re.match(r"^\[\d+\]\s", line.strip()):
            _add_body_no_indent(doc, line.strip())
            i += 1
            continue

        italic = line.startswith("*") and line.endswith("*") and not line.startswith("**")
        body_line = line
        if italic and len(line) >= 2:
            body_line = line[1:-1]
            _add_body_no_indent(doc, body_line, italic=True)
        else:
            if line.startswith("**关键词：**"):
                _add_body_no_indent(doc, line)
            else:
                _add_body_paragraph(doc, line)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Written: {out_path}")


def main() -> None:
    md = MD_PATH if len(sys.argv) < 2 else Path(sys.argv[1])
    out = OUT_PATH if len(sys.argv) < 3 else Path(sys.argv[2])
    if not md.is_file():
        raise SystemExit(f"Markdown not found: {md}")
    convert(md, out)


if __name__ == "__main__":
    main()
